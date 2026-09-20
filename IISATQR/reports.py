"""Excel and PDF report builders for IISAT QR.

Report generation is intentionally isolated from route handling so templates,
API endpoints, and future scheduled reports can reuse the same output logic.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

BLACK = "111111"
DARK = "202124"
LIGHT = "F4F4F4"
GREEN = "197A55"
RED = "B23A3A"


# ---------------------------------------------------------------------------
# Shared formatting helpers
# ---------------------------------------------------------------------------

def _local(dt):
    if not dt:
        return "—"

    try:
        return (
            dt.astimezone().strftime("%d %b %Y, %I:%M %p")
            if dt.tzinfo
            else dt.strftime("%d %b %Y, %I:%M %p")
        )
    except Exception:
        return str(dt)


def _rows(session):
    records = sorted(
        session.records,
        key=lambda r: r.student.registration_no
    )

    return [
        (
            r.student.registration_no,
            r.student.name,
            r.status,
            _local(r.marked_at),
        )
        for r in records
    ]


# ---------------------------------------------------------------------------
# Single-session Excel report
# ---------------------------------------------------------------------------

def build_excel(session):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    course = session.course
    teacher = session.creator

    total = len(session.records)
    present = sum(
        1
        for r in session.records
        if r.status == "PRESENT"
    )
    absent = total - present

    ws.merge_cells("A1:F1")
    ws["A1"] = "IISAT UNIVERSITY, GUJRANWALA — ATTENDANCE REPORT"

    ws["A1"].font = Font(
        size=16,
        bold=True,
        color="FFFFFF"
    )

    ws["A1"].fill = PatternFill(
        "solid",
        fgColor=BLACK
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws.row_dimensions[1].height = 30

    details = [
        (
            "Subject",
            f"{course.code} — {course.name}",
            "Teacher",
            teacher.name,
        ),
        (
            "Department",
            course.department,
            "Semester / Section",
            f"{course.semester} / {course.section}",
        ),
        (
            "Topic",
            session.topic or "Regular Class",
            "Room",
            course.room or "—",
        ),
        (
            "Started",
            _local(session.started_at),
            "Ended",
            _local(session.ended_at),
        ),
        (
            "Class Strength",
            total,
            "Present / Absent",
            f"{present} / {absent}",
        ),
    ]

    r = 3

    for a, b, c, d in details:
        ws[f"A{r}"] = a
        ws[f"B{r}"] = b
        ws.merge_cells(
            start_row=r,
            start_column=2,
            end_row=r,
            end_column=3
        )

        ws[f"D{r}"] = c
        ws[f"E{r}"] = d
        ws.merge_cells(
            start_row=r,
            start_column=5,
            end_row=r,
            end_column=6
        )

        ws[f"A{r}"].font = Font(bold=True)
        ws[f"D{r}"].font = Font(bold=True)

        r += 1

    r += 1

    headers = [
        "Sr.",
        "Registration No.",
        "Student Name",
        "Status",
        "Marked At",
        "Verification",
    ]

    for col, title in enumerate(headers, 1):
        cell = ws.cell(r, col, title)

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=BLACK
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    thin = Side(
        style="thin",
        color="D9D9D9"
    )

    for idx, (reg, name, status, marked) in enumerate(
        _rows(session),
        1
    ):
        row = r + idx

        values = [
            idx,
            reg,
            name,
            status.title(),
            marked,
            "QR Verified" if status == "PRESENT" else "Roster",
        ]

        for col, value in enumerate(values, 1):
            c = ws.cell(row, col, value)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="center")

        ws.cell(row, 4).font = Font(
            bold=True,
            color=GREEN if status == "PRESENT" else RED
        )

        if idx % 2 == 0:
            for col in range(1, 7):
                ws.cell(row, col).fill = PatternFill(
                    "solid",
                    fgColor=LIGHT
                )

    for i, w in enumerate(
        [8, 22, 30, 14, 25, 18],
        1
    ):
        ws.column_dimensions[
            get_column_letter(i)
        ].width = w

    ws.freeze_panes = f"A{r + 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    ws.oddFooter.center.text = "Generated by IISAT QR"

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    return out


# ---------------------------------------------------------------------------
# Single-session PDF report
# ---------------------------------------------------------------------------

def build_pdf(session):
    out = BytesIO()

    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title = ParagraphStyle(
        "title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        textColor=colors.HexColor("#111111"),
        alignment=TA_CENTER,
    )

    sub = ParagraphStyle(
        "sub",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#666666"),
    )

    course = session.course
    teacher = session.creator

    total = len(session.records)
    present = sum(
        1
        for r in session.records
        if r.status == "PRESENT"
    )
    absent = total - present

    story = [
        Paragraph(
            "IISAT UNIVERSITY, GUJRANWALA",
            title
        ),
        Paragraph(
            "IISAT QR — Smart Attendance Report",
            sub
        ),
        Spacer(1, 4 * mm),
    ]

    info = [
        [
            "Subject",
            f"{course.code} — {course.name}",
            "Teacher",
            teacher.name,
        ],
        [
            "Department",
            course.department,
            "Semester / Section",
            f"{course.semester} / {course.section}",
        ],
        [
            "Topic",
            session.topic or "Regular Class",
            "Room",
            course.room or "—",
        ],
        [
            "Started",
            _local(session.started_at),
            "Ended",
            _local(session.ended_at),
        ],
        [
            "Strength",
            str(total),
            "Present / Absent",
            f"{present} / {absent}",
        ],
    ]

    t = Table(
        info,
        colWidths=[
            27 * mm,
            68 * mm,
            32 * mm,
            82 * mm,
        ]
    )

    t.setStyle(
        TableStyle([
            (
                "FONTNAME",
                (0, 0),
                (-1, -1),
                "Helvetica",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8.5,
            ),
            (
                "FONTNAME",
                (0, 0),
                (0, -1),
                "Helvetica-Bold",
            ),
            (
                "FONTNAME",
                (2, 0),
                (2, -1),
                "Helvetica-Bold",
            ),
            (
                "BACKGROUND",
                (0, 0),
                (-1, -1),
                colors.HexColor("#F7F7F7"),
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D6D6D6"),
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    story += [
        t,
        Spacer(1, 4 * mm)
    ]

    data = [[
        "Sr.",
        "Registration No.",
        "Student Name",
        "Status",
        "Marked At",
    ]]

    for idx, (
        reg,
        name,
        status,
        marked
    ) in enumerate(_rows(session), 1):
        data.append([
            idx,
            reg,
            name,
            status.title(),
            marked,
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            14 * mm,
            45 * mm,
            78 * mm,
            28 * mm,
            52 * mm,
        ]
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#111111"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8.5,
            ),
            (
                "ALIGN",
                (0, 0),
                (0, -1),
                "CENTER",
            ),
            (
                "ALIGN",
                (3, 1),
                (3, -1),
                "CENTER",
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F5F5F5"),
                ],
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D6D6D6"),
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    story.append(table)

    doc.build(story)
    out.seek(0)

    return out


# ---------------------------------------------------------------------------
# Course analytics PDF report
# ---------------------------------------------------------------------------

def build_analytics_pdf(
    course,
    student_stats,
    overall_percentage,
    total_sessions,
    teacher=None,
    month=None,
    year=None,
):
    """
    Build a full-class attendance analytics PDF.

    student_stats is expected to contain dictionaries with:
    student, total, present, absent, percentage.
    """

    out = BytesIO()

    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "analytics_title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        textColor=colors.HexColor("#111111"),
        alignment=TA_CENTER,
    )

    subtitle_style = ParagraphStyle(
        "analytics_subtitle",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#666666"),
    )

    small_style = ParagraphStyle(
        "analytics_small",
        parent=styles["Normal"],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#333333"),
    )

    month_names = {
        1: "January",
        2: "February",
        3: "March",
        4: "April",
        5: "May",
        6: "June",
        7: "July",
        8: "August",
        9: "September",
        10: "October",
        11: "November",
        12: "December",
    }

    if month and year:
        period_label = f"{month_names.get(month, month)} {year}"
    elif month:
        period_label = month_names.get(month, str(month))
    elif year:
        period_label = str(year)
    else:
        period_label = "All Time"

    teacher_name = (
        teacher.name
        if teacher is not None
        else "—"
    )

    total_students = len(student_stats)

    total_present = sum(
        row["present"]
        for row in student_stats
    )

    total_absent = sum(
        row["absent"]
        for row in student_stats
    )

    story = [
        Paragraph(
            "IISAT UNIVERSITY, GUJRANWALA",
            title_style
        ),
        Paragraph(
            "IISAT QR — Course Attendance Analytics Report",
            subtitle_style
        ),
        Spacer(1, 4 * mm),
    ]

    info = [
        [
            "Course",
            f"{course.code} — {course.name}",
            "Generated By",
            teacher_name,
        ],
        [
            "Department",
            course.department,
            "Semester / Section",
            f"{course.semester} / {course.section}",
        ],
        [
            "Room",
            course.room or "—",
            "Attendance Period",
            period_label,
        ],
        [
            "Students",
            str(total_students),
            "Total Sessions",
            str(total_sessions),
        ],
        [
            "Overall Attendance",
            f"{overall_percentage:.1f}%",
            "Present / Absent",
            f"{total_present} / {total_absent}",
        ],
    ]

    info_table = Table(
        info,
        colWidths=[
            31 * mm,
            70 * mm,
            37 * mm,
            72 * mm,
        ]
    )

    info_table.setStyle(
        TableStyle([
            (
                "FONTNAME",
                (0, 0),
                (-1, -1),
                "Helvetica",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8.5,
            ),
            (
                "FONTNAME",
                (0, 0),
                (0, -1),
                "Helvetica-Bold",
            ),
            (
                "FONTNAME",
                (2, 0),
                (2, -1),
                "Helvetica-Bold",
            ),
            (
                "BACKGROUND",
                (0, 0),
                (-1, -1),
                colors.HexColor("#F7F7F7"),
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D6D6D6"),
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    story.extend([
        info_table,
        Spacer(1, 5 * mm),
        Paragraph(
            "Student Attendance Summary",
            small_style
        ),
        Spacer(1, 2 * mm),
    ])

    data = [[
        "Sr.",
        "Registration No.",
        "Student Name",
        "Total Classes",
        "Present",
        "Absent",
        "Attendance %",
    ]]

    for index, row in enumerate(
        student_stats,
        start=1
    ):
        data.append([
            index,
            row["student"].registration_no,
            row["student"].name,
            row["total"],
            row["present"],
            row["absent"],
            f'{row["percentage"]:.1f}%',
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            12 * mm,
            42 * mm,
            66 * mm,
            28 * mm,
            22 * mm,
            22 * mm,
            30 * mm,
        ]
    )

    style_commands = [
        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#111111"),
        ),
        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white,
        ),
        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold",
        ),
        (
            "FONTNAME",
            (0, 1),
            (-1, -1),
            "Helvetica",
        ),
        (
            "FONTSIZE",
            (0, 0),
            (-1, -1),
            8.3,
        ),
        (
            "ALIGN",
            (0, 0),
            (0, -1),
            "CENTER",
        ),
        (
            "ALIGN",
            (3, 1),
            (-1, -1),
            "CENTER",
        ),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor("#F5F5F5"),
            ],
        ),
        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.35,
            colors.HexColor("#D6D6D6"),
        ),
        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),
        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),
    ]

    # Color attendance percentages by simple visual bands.
    for row_index, row in enumerate(
        student_stats,
        start=1
    ):
        percentage = row["percentage"]

        if percentage >= 75:
            percentage_color = colors.HexColor("#197A55")
        elif percentage >= 60:
            percentage_color = colors.HexColor("#B07A12")
        else:
            percentage_color = colors.HexColor("#B23A3A")

        style_commands.append(
            (
                "TEXTCOLOR",
                (6, row_index),
                (6, row_index),
                percentage_color,
            )
        )

        style_commands.append(
            (
                "FONTNAME",
                (6, row_index),
                (6, row_index),
                "Helvetica-Bold",
            )
        )

    table.setStyle(
        TableStyle(style_commands)
    )

    story.append(table)

    doc.build(story)
    out.seek(0)

    return out
