
"""Course analytics pages and downloadable analytics reports."""

from io import BytesIO

from flask import flash, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import bp
from .helpers import can_manage_course, current_teacher, teacher_required
from ..extensions import db
from ..models import Course
from ..reports import build_analytics_pdf
from ..services.analytics import build_course_analytics


@bp.get("/courses/<int:course_id>/analytics")
@teacher_required
def course_analytics(course_id):
    course = db.session.get(Course, course_id)

    if not course:
        return "Course not found", 404

    if not can_manage_course(
        current_teacher(),
        course
    ):
        return "Not authorized", 403

    month = request.args.get(
        "month",
        type=int
    )

    year = request.args.get(
        "year",
        type=int
    )

    if month is not None and not 1 <= month <= 12:
        flash(
            "Please select a valid month.",
            "warning"
        )
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    if year is not None and not 2000 <= year <= 2100:
        flash(
            "Please enter a valid year.",
            "warning"
        )
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    analytics = build_course_analytics(
        course,
        month=month,
        year=year
    )

    return render_template(
        "course_analytics.html",
        teacher=current_teacher(),
        course=course,
        month=month,
        year=year,
        **analytics
    )


@bp.get("/courses/<int:course_id>/analytics/export/excel")
@teacher_required
def export_analytics_excel(course_id):
    course = db.session.get(Course, course_id)

    if not course:
        return "Course not found", 404

    if not can_manage_course(current_teacher(), course):
        return "Not authorized", 403

    month = request.args.get("month", type=int)
    year = request.args.get("year", type=int)

    if month is not None and not 1 <= month <= 12:
        flash("Please select a valid month.", "warning")
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    if year is not None and not 2000 <= year <= 2100:
        flash("Please enter a valid year.", "warning")
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    analytics = build_course_analytics(
        course,
        month=month,
        year=year
    )

    student_stats = analytics["student_stats"]
    overall_percentage = analytics["overall_percentage"]
    total_sessions = analytics["total_sessions"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance Analytics"

    # Report title
    sheet.merge_cells("A1:F1")
    title_cell = sheet["A1"]
    title_cell.value = "IISAT QR Attendance Analytics Report"
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    # Course details
    details = [
        ("Course", f"{course.code} - {course.name}"),
        ("Department", course.department),
        ("Semester", course.semester),
        ("Section", course.section),
        ("Room", course.room or "Not Assigned"),
        ("Month", month or "All Months"),
        ("Year", year or "All Years"),
        ("Total Sessions", total_sessions),
        ("Overall Attendance", f"{overall_percentage}%"),
    ]

    start_detail_row = 3

    for index, (label, value) in enumerate(
        details,
        start=start_detail_row
    ):
        sheet.cell(
            row=index,
            column=1,
            value=label
        ).font = Font(bold=True)

        sheet.cell(
            row=index,
            column=2,
            value=value
        )

    header_row = start_detail_row + len(details) + 2

    headers = [
        "Registration No",
        "Student Name",
        "Total Classes",
        "Present",
        "Absent",
        "Attendance %"
    ]

    for column, heading in enumerate(headers, start=1):
        cell = sheet.cell(
            row=header_row,
            column=column,
            value=heading
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="111111"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    current_row = header_row + 1

    for row in student_stats:
        sheet.cell(
            row=current_row,
            column=1,
            value=row["student"].registration_no
        )

        sheet.cell(
            row=current_row,
            column=2,
            value=row["student"].name
        )

        sheet.cell(
            row=current_row,
            column=3,
            value=row["total"]
        )

        sheet.cell(
            row=current_row,
            column=4,
            value=row["present"]
        )

        sheet.cell(
            row=current_row,
            column=5,
            value=row["absent"]
        )

        percentage_cell = sheet.cell(
            row=current_row,
            column=6,
            value=row["percentage"] / 100
        )

        percentage_cell.number_format = "0.0%"

        current_row += 1

    # Professional column widths
    widths = {
        "A": 22,
        "B": 30,
        "C": 15,
        "D": 12,
        "E": 12,
        "F": 18,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # Center numeric columns
    if current_row > header_row + 1:
        for row in sheet.iter_rows(
            min_row=header_row + 1,
            max_row=current_row - 1,
            min_col=3,
            max_col=6
        ):
            for cell in row:
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

    sheet.freeze_panes = f"A{header_row + 1}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    period_parts = []

    if month:
        period_parts.append(f"M{month:02d}")

    if year:
        period_parts.append(str(year))

    period_suffix = (
        "_" + "_".join(period_parts)
        if period_parts
        else "_All_Time"
    )

    filename = (
        f"IISAT_QR_{course.code}"
        f"{period_suffix}_Attendance_Analytics.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


@bp.get("/courses/<int:course_id>/analytics/export/pdf")
@teacher_required
def export_analytics_pdf(course_id):
    course = db.session.get(Course, course_id)

    if not course:
        return "Course not found", 404

    teacher = current_teacher()

    if not can_manage_course(teacher, course):
        return "Not authorized", 403

    month = request.args.get("month", type=int)
    year = request.args.get("year", type=int)

    if month is not None and not 1 <= month <= 12:
        flash("Please select a valid month.", "warning")
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    if year is not None and not 2000 <= year <= 2100:
        flash("Please enter a valid year.", "warning")
        return redirect(
            url_for(
                "main.course_analytics",
                course_id=course.id
            )
        )

    analytics = build_course_analytics(
        course,
        month=month,
        year=year
    )

    pdf_file = build_analytics_pdf(
        course=course,
        student_stats=analytics["student_stats"],
        overall_percentage=analytics["overall_percentage"],
        total_sessions=analytics["total_sessions"],
        teacher=teacher,
        month=month,
        year=year,
    )

    period_parts = []

    if month:
        period_parts.append(f"M{month:02d}")

    if year:
        period_parts.append(str(year))

    period_suffix = (
        "_" + "_".join(period_parts)
        if period_parts
        else "_All_Time"
    )

    filename = (
        f"IISAT_QR_{course.code}"
        f"{period_suffix}_Attendance_Analytics.pdf"
    )

    return send_file(
        pdf_file,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf"
    )


